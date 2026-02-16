#!/usr/bin/env python3
import argparse
import csv
import datetime as dt
import json
import pathlib
import re
import subprocess
from collections import defaultdict
from typing import Dict, List, Tuple, Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


def run(cmd: List[str], cwd: pathlib.Path) -> str:
    return subprocess.check_output(cmd, cwd=str(cwd), text=True)


def git_commits(repo: pathlib.Path, start_date: str, end_date: str) -> List[Dict[str, Any]]:
    log = run(
        [
            "git",
            "log",
            "--reverse",
            "--pretty=format:%H|%cI",
            f"--since={start_date}T00:00:00Z",
            f"--until={end_date}T23:59:59Z",
        ],
        repo,
    )
    commits = []
    for line in log.splitlines():
        commit, iso = line.split("|", 1)
        timestamp = dt.datetime.fromisoformat(iso.replace("Z", "+00:00"))
        commits.append({"commit": commit, "date": timestamp.date().isoformat(), "datetime": timestamp})
    return commits


def parse_ir_snapshots(ir_dir: pathlib.Path, commits: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    snapshots = []
    commit_map = {c["commit"]: c for c in commits}
    for path in sorted(ir_dir.glob("IR*_*.json")):
        m = re.match(r"IR(\d+)_([0-9a-fA-F]+)\.json", path.name)
        if not m:
            continue
        idx = int(m.group(1))
        short = m.group(2).lower()
        commit = next((c for c in commit_map if c.startswith(short)), None)
        if not commit:
            continue
        with path.open("r", encoding="utf-8") as fh:
            data = json.load(fh)
        snapshots.append(
            {
                "index": idx,
                "path": path,
                "short": short,
                "commit": commit,
                "date": commit_map[commit]["date"],
                "datetime": commit_map[commit]["datetime"],
                "ir": data,
            }
        )
    snapshots.sort(key=lambda s: (s["index"], s["datetime"]))
    return snapshots


def iso_week_label(date_str: str) -> str:
    y, w, _ = dt.date.fromisoformat(date_str).isocalendar()
    return f"{y}-W{w:02d}"


def extract_service_paths(snapshot: Dict[str, Any]) -> List[Tuple[str, str]]:
    rules = []
    for ms in snapshot["ir"].get("microservices", []) or []:
        name = ms.get("name")
        path = (ms.get("path") or "").strip("/")
        if name and path:
            rules.append((name, path))
    rules.sort(key=lambda item: len(item[1]), reverse=True)
    return rules


def map_file_to_service(file_path: str, rules: List[Tuple[str, str]]) -> str:
    normalized = file_path.strip("/")
    for service, prefix in rules:
        if normalized == prefix or normalized.startswith(prefix + "/"):
            return service
    if "/" in normalized:
        return normalized.split("/", 1)[0]
    return "ROOT"


def parse_numstat(repo: pathlib.Path, commit: str) -> List[Dict[str, Any]]:
    out = run(["git", "show", "--numstat", "--format=", commit], repo)
    rows = []
    for line in out.splitlines():
        parts = line.split("\t")
        if len(parts) != 3:
            continue
        add, delete, file_path = parts
        add_n = 0 if add == "-" else int(add)
        del_n = 0 if delete == "-" else int(delete)
        rows.append({
            "file": file_path,
            "loc_added": add_n,
            "loc_deleted": del_n,
            "loc_delta": add_n - del_n,
        })
    return rows


def list_service_classes(ms: Dict[str, Any]) -> List[Dict[str, Any]]:
    classes = []
    for key in ["controllers", "services", "repositories", "entities", "feignClients"]:
        classes.extend(ms.get(key, []) or [])
    return classes


def operation_from_method(service_name: str, method: Dict[str, Any]) -> Dict[str, Any]:
    params = [p.get("type") for p in (method.get("parameters") or []) if p.get("type")]
    anns = [f"{a.get('name')} - {a.get('contents')}" for a in (method.get("annotations") or []) if a.get("name")]
    return {
        "name": f"{service_name}::{method.get('name', '')}",
        "responseType": method.get("returnType") or "",
        "params": params,
        "usingTypes": anns,
    }


def lomlc(operations: List[Dict[str, Any]]) -> float:
    if len(operations) <= 1:
        return 1.0
    pairs = 0
    acc = 0.0
    for i in range(len(operations)):
        for j in range(i + 1, len(operations)):
            op1, op2 = operations[i], operations[j]
            p1, p2 = set(op1["params"]), set(op2["params"])
            union_size = len(p1 | p2)
            ids = 1.0 if union_size == 0 else len(p1 & p2) / union_size
            ods = 1.0 if op1["responseType"] == op2["responseType"] else 0.0
            acc += 1 - ((ids + ods) / 2)
            pairs += 1
    return acc / pairs if pairs else 1.0


def endpoints_by_service(snapshot: Dict[str, Any]) -> Dict[str, List[Tuple[str, str]]]:
    data = defaultdict(list)
    for ms in snapshot["ir"].get("microservices", []) or []:
        sname = ms.get("name")
        for controller in ms.get("controllers", []) or []:
            for ep in controller.get("endpoints", []) or []:
                method = (ep.get("httpMethod") or "ALL").upper()
                url = (ep.get("url") or "").split("?", 1)[0]
                if sname and url:
                    data[sname].append((method, url))
    return data


def match_restcall_target(source: str, method: str, url: str, endpoint_index: Dict[str, List[Tuple[str, str]]]) -> List[str]:
    clean = url.replace("{?}", "").split("?", 1)[0]
    matches = []
    for svc, eps in endpoint_index.items():
        if svc == source:
            continue
        for ep_method, ep_url in eps:
            if clean == ep_url and (method == ep_method or ep_method == "ALL"):
                matches.append(svc)
                break
    return matches


def compute_weekly_tables(snapshots: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    by_week = {}
    for snap in snapshots:
        week = iso_week_label(snap["date"])
        by_week[week] = snap

    service_rows = []
    pair_rows = []
    previous_pair_scores: Dict[Tuple[str, str], int] = {}

    for week in sorted(by_week):
        snap = by_week[week]
        endpoint_index = endpoints_by_service(snap)

        pair_scores = defaultdict(int)
        for ms in snap["ir"].get("microservices", []) or []:
            sname = ms.get("name")
            if not sname:
                continue

            classes = list_service_classes(ms)
            methods = [m for c in classes for m in (c.get("methods") or [])]
            controllers = ms.get("controllers", []) or []
            operations = [operation_from_method(sname, m) for c in controllers for m in (c.get("methods") or [])]
            dep_targets = set()

            for c in classes:
                for mc in c.get("methodCalls", []) or []:
                    tgt = mc.get("objectType")
                    if tgt:
                        dep_targets.add(tgt)
                for rc in c.get("restCalls", []) or []:
                    rmethod = (rc.get("httpMethod") or "GET").upper()
                    rurl = rc.get("url") or ""
                    for target in match_restcall_target(sname, rmethod, rurl, endpoint_index):
                        key = tuple(sorted((sname, target)))
                        pair_scores[key] += 1

            service_rows.append(
                {
                    "week": week,
                    "commit": snap["commit"],
                    "service": sname,
                    "LOMLC": round(lomlc(operations), 6),
                    "internal_dependency_count": len(dep_targets),
                    "function_count": len(methods),
                }
            )

        all_pairs = set(previous_pair_scores) | set(pair_scores)
        for pair in sorted(all_pairs):
            score = pair_scores.get(pair, 0)
            prev = previous_pair_scores.get(pair, 0)
            pair_rows.append(
                {
                    "week": week,
                    "commit": snap["commit"],
                    "service_1": pair[0],
                    "service_2": pair[1],
                    "SSIC": score,
                    "SSIC_delta": score - prev,
                }
            )
        previous_pair_scores = dict(pair_scores)

    return service_rows, pair_rows


def write_csv(path: pathlib.Path, rows: List[Dict[str, Any]], headers: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in headers})


def enrich_workbook(workbook_path: pathlib.Path, trace_rows, service_rows, pair_rows) -> None:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to enrich workbook.")

    wb = load_workbook(workbook_path)

    definitions = [
        (
            "Commit-Service-Trace",
            ["commit", "date", "file", "service", "loc_added", "loc_deleted", "loc_delta"],
            trace_rows,
        ),
        (
            "Service-Weekly-Metrics",
            ["week", "commit", "service", "LOMLC", "internal_dependency_count", "function_count"],
            service_rows,
        ),
        (
            "Service-Pair-Coupling",
            ["week", "commit", "service_1", "service_2", "SSIC", "SSIC_delta"],
            pair_rows,
        ),
    ]

    for sheet_name, headers, rows in definitions:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])

    wb.save(workbook_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extend Spinnaker metrics output with service trace/coupling tables.")
    parser.add_argument("--repo", required=True, type=pathlib.Path)
    parser.add_argument("--ir-dir", required=True, type=pathlib.Path)
    parser.add_argument("--workbook", required=True, type=pathlib.Path)
    parser.add_argument("--output-dir", required=True, type=pathlib.Path)
    parser.add_argument("--start-date", required=True)
    parser.add_argument("--end-date", required=True)
    args = parser.parse_args()

    commits = git_commits(args.repo, args.start_date, args.end_date)
    snapshots = parse_ir_snapshots(args.ir_dir, commits)
    if not commits:
        raise SystemExit("No commits found for the requested date window.")
    if not snapshots:
        raise SystemExit("No IR snapshots found to build weekly service metrics.")

    snapshot_rules = sorted([(snap["datetime"], extract_service_paths(snap)) for snap in snapshots], key=lambda t: t[0])

    trace_rows = []
    sorted_commits = sorted(commits, key=lambda c: c["datetime"])
    for commit in sorted_commits:
        rules = snapshot_rules[0][1]
        for snap_time, snap_rules in snapshot_rules:
            if snap_time <= commit["datetime"]:
                rules = snap_rules
            else:
                break
        for numstat in parse_numstat(args.repo, commit["commit"]):
            trace_rows.append(
                {
                    "commit": commit["commit"],
                    "date": commit["date"],
                    "file": numstat["file"],
                    "service": map_file_to_service(numstat["file"], rules),
                    "loc_added": numstat["loc_added"],
                    "loc_deleted": numstat["loc_deleted"],
                    "loc_delta": numstat["loc_delta"],
                }
            )

    service_rows, pair_rows = compute_weekly_tables(snapshots)

    out = args.output_dir
    write_csv(out / "commit_service_trace.csv", trace_rows, ["commit", "date", "file", "service", "loc_added", "loc_deleted", "loc_delta"])
    write_csv(out / "service_weekly_metrics.csv", service_rows, ["week", "commit", "service", "LOMLC", "internal_dependency_count", "function_count"])
    write_csv(out / "service_pair_coupling.csv", pair_rows, ["week", "commit", "service_1", "service_2", "SSIC", "SSIC_delta"])

    summary = {
        "commits": len(commits),
        "trace_rows": len(trace_rows),
        "weeks": len({row["week"] for row in service_rows}),
        "service_rows": len(service_rows),
        "pair_rows": len(pair_rows),
    }
    with (out / "extended_metrics_summary.json").open("w", encoding="utf-8") as fh:
        json.dump(summary, fh, indent=2)

    enrich_workbook(args.workbook, trace_rows, service_rows, pair_rows)


if __name__ == "__main__":
    main()
